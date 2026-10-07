from pathlib import Path

import pytest # type: ignore
from fastapi import HTTPException

from backend import main
from backend.tools import document_tool


#Tests to ensure that the document generation functions correctly handle Unicode content and produce unique filenames for each generated document. The tests use monkeypatching to set the GENERATED_DIR to a temporary path, generate documents in different formats (PDF, DOCX, XLSX), and verify that the files are created with unique names and correct download URLs.
@pytest.mark.parametrize("file_type", ["pdf", "docx", "xlsx"])
def test_documents_support_unicode_and_unique_names(tmp_path, monkeypatch, file_type):
    monkeypatch.setattr(document_tool, "GENERATED_DIR", tmp_path)
    content = "**Report**\nResearch — result • Working 😊"

    first = document_tool.generate_document(content, file_type)
    second = document_tool.generate_document(content, file_type)

    assert first["filename"] != second["filename"]
    assert (tmp_path / first["filename"]).is_file()
    assert first["download_url"] == f"/download/{first['filename']}"

#Test to ensure that the XLSX document generation function correctly skips the markdown alignment row when creating the spreadsheet. The test uses monkeypatching to set the GENERATED_DIR to a temporary path, generates an XLSX document from markdown content with an alignment row, and verifies that the resulting spreadsheet has the expected number of rows and correct cell values.
def test_xlsx_skips_markdown_alignment_row(tmp_path, monkeypatch):
    monkeypatch.setattr(document_tool, "GENERATED_DIR", tmp_path)
    result = document_tool.generate_document(
        "| Name | Value |\n| :--- | ---: |\n| A | 1 |",
        "xlsx",
    )

    from openpyxl import load_workbook

    sheet = load_workbook(tmp_path / result["filename"]).active
    assert sheet.max_row == 2
    assert sheet["A2"].value == "A"


#Test to ensure that the download_file endpoint correctly handles requests for missing files by returning a 404 HTTPException. The test uses monkeypatching to set the GENERATED_DIR to a temporary path and attempts to download a non-existent file, verifying that the appropriate exception is raised with the correct status code.
def test_download_returns_404_for_missing_file(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "GENERATED_DIR", tmp_path)

    with pytest.raises(HTTPException) as exc_info:
        main.download_file("missing.pdf")

    assert exc_info.value.status_code == 404


#test to ensure that the download_file endpoint correctly serves files from the GENERATED_DIR using only the filename, not an absolute server path. The test creates a temporary file in the GENERATED_DIR, calls the download_file function with the filename, and verifies that the response path matches the generated file's path, confirming that the endpoint serves files securely without exposing server paths.
def test_download_uses_filename_not_absolute_server_path(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "GENERATED_DIR", tmp_path)
    generated = tmp_path / "result.pdf"
    generated.write_bytes(b"pdf")

    response = main.download_file("result.pdf")

    assert Path(response.path) == generated
