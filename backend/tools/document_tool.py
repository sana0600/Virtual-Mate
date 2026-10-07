import os
import re
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

try:
    from ..config import GENERATED_DIR
    from ..errors import DocumentGenerationError
except ImportError:
    from config import GENERATED_DIR
    from errors import DocumentGenerationError


SUPPORTED_TYPES = {"pdf", "docx", "xlsx"}
FONT_CANDIDATES = (
    Path(os.getenv("VIRTUALMATE_PDF_FONT", "")),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
)
FONT_BOLD_CANDIDATES = (
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
)


def _bold_segments(text: str) -> list[tuple[str, bool]]:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    return [
        (part[2:-2], True) if part.startswith("**") and part.endswith("**") else (part, False)
        for part in parts
        if part
    ]


def _register_pdf_fonts() -> tuple[str, str]:
    regular = next((path for path in FONT_CANDIDATES if path and path.is_file()), None)
    bold = next((path for path in FONT_BOLD_CANDIDATES if path.is_file()), None)
    if regular:
        pdfmetrics.registerFont(TTFont("VirtualMateSans", str(regular)))
        if bold:
            pdfmetrics.registerFont(TTFont("VirtualMateSans-Bold", str(bold)))
        else:
            pdfmetrics.registerFont(TTFont("VirtualMateSans-Bold", str(regular)))
        return "VirtualMateSans", "VirtualMateSans-Bold"
    return "Helvetica", "Helvetica-Bold"


def _pdf_markup(line: str, unicode_font: bool) -> str:
    if not unicode_font:
        line = line.encode("latin-1", errors="replace").decode("latin-1")
    return "".join(
        f"<b>{escape(text)}</b>" if is_bold else escape(text)
        for text, is_bold in _bold_segments(line)
    )


def _generate_pdf(content: str, filename: Path) -> None:
    regular_font, bold_font = _register_pdf_fonts()
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "VirtualMateBody",
        parent=styles["BodyText"],
        fontName=regular_font,
        fontSize=10.5,
        leading=15,
        alignment=TA_LEFT,
        spaceAfter=5,
    )
    pdfmetrics.registerFontFamily(
        "VirtualMateSans" if regular_font == "VirtualMateSans" else "Helvetica",
        normal=regular_font,
        bold=bold_font,
    )
    story = []
    for line in str(content).splitlines() or [""]:
        if not line.strip():
            story.append(Spacer(1, 5 * mm))
            continue
        story.append(Paragraph(_pdf_markup(line, regular_font == "VirtualMateSans"), body))

    document = SimpleDocTemplate(
        str(filename),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="VirtualMate Generated Report",
    )
    document.build(story)


def _generate_docx(content: str, filename: Path) -> None:
    document = Document()
    for line in str(content).splitlines():
        paragraph = document.add_paragraph()
        for text, is_bold in _bold_segments(line):
            paragraph.add_run(text).bold = is_bold
    document.save(filename)


def _generate_xlsx(content: str, filename: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "VirtualMate Report"
    row_number = 1

    for raw_line in str(content).splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("|") and line.endswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            for column, value in enumerate(cells, start=1):
                cell = worksheet.cell(row=row_number, column=column, value=value)
                if row_number == 1:
                    cell.font = Font(bold=True)
        else:
            worksheet.cell(row=row_number, column=1, value=line)
        row_number += 1
    workbook.save(filename)


def generate_document(content: str, file_type: str) -> dict[str, str]:
    file_type = file_type.lower()
    if file_type not in SUPPORTED_TYPES:
        raise DocumentGenerationError(f"Unsupported document type: {file_type}")

    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    filename = GENERATED_DIR / f"virtualmate_{uuid4().hex[:12]}.{file_type}"

    try:
        if file_type == "pdf":
            _generate_pdf(content, filename)
        elif file_type == "docx":
            _generate_docx(content, filename)
        else:
            _generate_xlsx(content, filename)
    except Exception as exc:
        filename.unlink(missing_ok=True)
        raise DocumentGenerationError(
            f"Could not generate the requested {file_type.upper()} file."
        ) from exc

    return {
        "message": f"{file_type.upper()} generated successfully",
        "filename": filename.name,
        "download_url": f"/download/{filename.name}",
    }
